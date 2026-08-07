# -*- coding: utf-8 -*-
"""Gera o relatorio Word do SE Mapper.

Uso: python gerar_docx.py <payload.json> <saida.docx>

O payload traz {tipo: 'mapeamento'|'comparacao', ...} montado pelo lib/relatorio.js.
"""
import json, sys, os, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

payload_path, saida = sys.argv[1], sys.argv[2]
with open(payload_path, encoding="utf-8") as f:
    D = json.load(f)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)
sec = doc.sections[0]
for lado in ("left_margin", "right_margin"):
    setattr(sec, lado, Cm(2))
sec.top_margin = sec.bottom_margin = Cm(2)
LARG = Cm(17)


def p(txt="", size=10, bold=False, italic=False, align=None, after=6):
    par = doc.add_paragraph()
    r = par.add_run(txt)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if align: par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    return par


def tabela(cab, linhas, larguras=None, fonte=8.5):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]; cel.text = ""
        r = cel.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(fonte)
    for ln in linhas:
        cells = t.add_row().cells
        for i, v in enumerate(ln):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            s = "" if v is None else str(v)
            neg = s.startswith("**") and s.endswith("**")
            if neg: s = s[2:-2]
            r = par.add_run(s); r.font.size = Pt(fonte); r.bold = neg
            par.paragraph_format.space_after = Pt(1)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def h(txt, nivel=1):
    doc.add_heading(txt, level=nivel)


def imagem(caminho, legenda):
    if not caminho or not os.path.exists(caminho):
        return
    doc.add_picture(caminho, width=LARG)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Figura — " + legenda); r.font.size = Pt(8); r.italic = True
    cap.paragraph_format.space_after = Pt(10)


# ============================================================ capa
tipo = D.get("tipo")
titulo = ("Relatório de Comparação de Ambientes" if tipo == "comparacao"
          else "Relatório de Mapeamento de Configuração")
p(titulo.upper(), size=19, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
p("SoftExpert Suite", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=22)

if tipo == "comparacao":
    o, d = D["origem"], D["destino"]
    tabela(["Item", "Valor"], [
        ["Ambiente de origem", "%s — %s" % (o["nome"], o["url"])],
        ["Ambiente de destino", "%s — %s" % (d["nome"], d["url"])],
        ["Comparação gerada em", D["geradoEm"]],
        ["Telas comparadas", D["resumo"]["telas"]],
        ["Método", "Automação Playwright sobre sessão autenticada manualmente. Coleta somente leitura."],
        ["Chave de comparação", "código da tela (o identificador de página muda entre ambientes)"],
    ], larguras=[Cm(5), Cm(12)], fonte=9)
else:
    a = D["ambiente"]; q = D["qualidade"]
    tabela(["Item", "Valor"], [
        ["Ambiente", a["nome"]],
        ["Endereço", a["url"]],
        ["Papel", a["papel"]],
        ["Coletado em", D.get("capturadoEm", "—")],
        ["Relatório gerado em", D["geradoEm"]],
        ["Método", "Automação Playwright sobre sessão autenticada manualmente. Coleta somente leitura."],
    ], larguras=[Cm(5), Cm(12)], fonte=9)

p()
p("ABRANGÊNCIA", size=10, bold=True)
p("O conteúdo reflete o que está visível para o perfil autenticado durante a coleta. A ausência de uma "
  "tela ou de um módulo não distingue item não licenciado de item sem permissão para aquele usuário.",
  size=9, italic=True)
doc.add_page_break()

# ============================================================ conteúdo
if tipo == "comparacao":
    r = D["resumo"]
    diverg = r["divergente"] + r["soOrigem"] + r["soDestino"]

    h("1. Resultado", 1)
    p(("%d de %d telas não conferem entre os ambientes." % (diverg, r["telas"])) if diverg
      else ("As %d telas conferem entre os ambientes." % r["telas"]), size=13, bold=True, after=10)
    tabela(["Indicador", "Quantidade"], [
        ["Telas comparadas", r["telas"]],
        ["Conferem", r["conforme"]],
        ["Contagem de registros diferente", r["divergente"]],
        ["Presentes só na origem", r["soOrigem"]],
        ["Presentes só no destino", r["soDestino"]],
        ["Não avaliadas em ao menos um lado", r["naoAvaliado"]],
        ["Componentes divergentes", r["componentesDivergentes"]],
    ], larguras=[Cm(11), Cm(6)], fonte=9)
    imagem(D.get("imagemTraco"), "Traço de divergência por tela")

    comp_dif = [c for c in D["componentes"] if c["status"] != "conforme"]
    if comp_dif:
        h("2. Componentes que não conferem", 1)
        tabela(["Componente", "Situação"], [[c["nome"], c["status"]] for c in comp_dif],
               larguras=[Cm(11), Cm(6)], fonte=9)

    h("3. Telas divergentes", 1)
    linhas = [[t["modulo"], t["codigo"] or "—", t["funcao"],
               "—" if t["qtdOrigem"] is None else t["qtdOrigem"],
               "—" if t["qtdDestino"] is None else t["qtdDestino"], t["status"]]
              for t in D["telas"] if t["status"] != "conforme"]
    if linhas:
        tabela(["Módulo", "Código", "Tela", "Origem", "Destino", "Situação"], linhas,
               larguras=[Cm(3), Cm(1.8), Cm(6), Cm(2), Cm(2), Cm(2.2)], fonte=8)
    else:
        p("Nenhuma divergência entre as telas comparadas.", italic=True)

    prof = [x for x in D.get("profundo", []) if x.get("divergentes")]
    if prof:
        h("4. Parâmetros divergentes nos registros", 1)
        for bloco in prof:
            p("%s · %s — %d de %d registros divergem" % (bloco["modulo"], bloco["codigo"],
              bloco["divergentes"], bloco["total"]), size=11, bold=True, after=4)
            linhas = []
            for reg in bloco["registros"]:
                if reg["status"] == "conforme":
                    continue
                difs = "; ".join("%s: %s → %s" % (d.get("rotulo"), d.get("origem"), d.get("destino"))
                                 for d in (reg.get("diferencas") or [])[:10]) or "—"
                linhas.append([reg["nome"], reg["status"], difs])
            tabela(["Registro", "Situação", "O que mudou"], linhas[:120],
                   larguras=[Cm(5), Cm(2.5), Cm(9.5)], fonte=8)

    h("5. Confiabilidade da coleta", 1)
    p("Uma varredura que falha muito não sustenta conclusão. Os indicadores abaixo separam falha da "
      "automação (exceção) de tela que abriu sem conteúdo legível (conferência manual).", size=9,
      italic=True, after=8)
    tabela(["Ambiente", "Telas", "Lidas", "Sem leitura", "Exceções", "Taxa de falha"],
           [[x["nome"], x["q"]["total"], x["q"]["lidas"], x["q"]["semConteudo"],
             x["q"]["comExcecao"], "%s%%" % x["q"]["taxaFalha"]] for x in D["qualidades"]],
           larguras=[Cm(4.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)], fonte=9)

else:
    q = D["qualidade"]
    h("1. Resultado da coleta", 1)
    p("%d de %d telas foram lidas com conteúdo." % (q["lidas"], q["total"]), size=13, bold=True, after=10)
    tabela(["Indicador", "Quantidade"], [
        ["Componentes", D["componentes"]],
        ["Telas de configuração previstas", D["previstas"]],
        ["Telas varridas", q["total"]],
        ["Telas lidas com conteúdo", "%d (%s%%)" % (q["lidas"], q["pctLidas"])],
        ["Telas sem conteúdo legível", "%d (%s%%)" % (q["semConteudo"], q["pctSemConteudo"])],
        ["Telas com exceção", "%d (%s%%)" % (q["comExcecao"], q["pctExcecao"])],
        ["Taxa de falha da coleta", "%s%%" % q["taxaFalha"]],
        ["Registros contados", D["registros"]],
    ], larguras=[Cm(11), Cm(6)], fonte=9)
    imagem(D.get("imagemTraco"), "Traço da coleta por tela")

    h("2. Telas por módulo", 1)
    tabela(["Módulo", "Telas", "Lidas", "Sem leitura", "Exceções", "Registros"],
           [[m["nome"], m["telas"], m["lidas"], m["semLeitura"], m["excecoes"], m["registros"]]
            for m in D["modulos"]],
           larguras=[Cm(5), Cm(2.4), Cm(2.4), Cm(2.6), Cm(2.2), Cm(2.4)], fonte=9)

    h("3. Inventário das telas de configuração", 1)
    tabela(["Módulo", "Código", "Tela", "Tipo", "Registros", "Situação"],
           [[t["modulo"], t["codigo"] or "—", t["funcao"], t.get("tipo") or "—",
             t.get("qtd", 0), t["situacao"]] for t in D["telas"]],
           larguras=[Cm(3), Cm(1.8), Cm(6.2), Cm(1.8), Cm(2), Cm(2.2)], fonte=7.5)

    pend = [t for t in D["telas"] if t["situacao"] != "lida"]
    h("4. Pendências para tratamento manual", 1)
    if pend:
        p("As telas abaixo abriram mas não entregaram conteúdo legível à automação, ou falharam com "
          "exceção. Precisam de conferência manual antes de qualquer conclusão.", size=9, italic=True, after=8)
        tabela(["Módulo", "Código", "Tela", "Motivo"],
               [[t["modulo"], t["codigo"] or "—", t["funcao"],
                 t.get("erro") or "a tela abriu, mas nada legível foi extraído"] for t in pend],
               larguras=[Cm(3), Cm(1.8), Cm(5.5), Cm(6.7)], fonte=8)
    else:
        p("Nenhuma pendência: todas as telas varridas renderam conteúdo.", italic=True)

    prof = D.get("profundo") or []
    if prof:
        h("5. Registros abertos em profundidade", 1)
        for bloco in prof:
            p("%s · %s — %d registro(s)" % (bloco["modulo"], bloco["codigo"], len(bloco["registros"])),
              size=11, bold=True, after=4)
            linhas = []
            for reg in bloco["registros"][:150]:
                ativos = sum(1 for c in reg.get("campos", []) if c.get("valor") == "ATIVO")
                linhas.append([reg["nome"], len(reg.get("campos", [])), ativos, reg.get("erro") or ""])
            tabela(["Registro", "Campos", "Ativos", "Observação"], linhas,
                   larguras=[Cm(7), Cm(2.2), Cm(2.2), Cm(5.6)], fonte=8)

doc.add_page_break()
h("Ressalvas da coleta", 1)
for txt in D.get("ressalvas", []):
    par = doc.add_paragraph(txt, style="List Bullet")
    par.runs[0].font.size = Pt(9.5)

p()
h("Revisão", 1)
tabela(["Revisado por", "Área", "Data", "Parecer"], [["", "", "", ""], ["", "", "", ""]],
       larguras=[Cm(4), Cm(4), Cm(2.5), Cm(6.5)], fonte=9)

doc.save(saida)
print("gerado: " + saida)
